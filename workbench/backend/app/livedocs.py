"""Live documents in replies (feat/v02-automations): editable docs, sheets (with formulas) and charts.

The assistant writes ```doc, ```sheet and ```chart blocks (auto_tools.LIVE_DOCS_PROMPT); the app renders them inline,
lets the user edit in place, and exports them here. Sheets keep their formulas in .xlsx; .csv and .pdf use the values
the app computed. Charts export as .png (drawn in the app), .csv, .xlsx (with a real Excel chart) and .pdf.
"""

from __future__ import annotations

import base64
import csv
import io
import json
import re

from fastapi import HTTPException

KINDS = ("doc", "sheet", "chart")
FORMATS = {"doc": ("docx", "pdf", "md"), "sheet": ("xlsx", "csv", "pdf"), "chart": ("png", "csv", "xlsx", "pdf")}
MIME = {"docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "pdf": "application/pdf",
        "csv": "text/csv; charset=utf-8", "md": "text/markdown; charset=utf-8", "png": "image/png"}


def safe_name(title: str) -> str:
    return re.sub(r"[^\w\- ()]+", "_", (title or "document").strip())[:80] or "document"


def _sheet(content) -> dict:
    d = json.loads(content) if isinstance(content, str) else dict(content or {})
    cols = [str(c) for c in d.get("columns") or []]
    rows = [list(r) for r in d.get("rows") or [] if isinstance(r, list)]
    if not cols and rows:
        cols, rows = [str(c) for c in rows[0]], rows[1:]
    return {"title": d.get("title") or "Sheet", "columns": cols, "rows": rows}


def _chart(content) -> dict:
    d = json.loads(content) if isinstance(content, str) else dict(content or {})
    labels = [str(x) for x in d.get("labels") or []]
    series = [{"name": str(s.get("name") or f"Series {i + 1}"), "values": [_num(v) for v in s.get("values") or []]}
              for i, s in enumerate(d.get("series") or []) if isinstance(s, dict)]
    return {"title": d.get("title") or "Chart", "type": d.get("type") if d.get("type") in ("bar", "line", "pie") else "bar",
            "labels": labels, "series": series}


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def _grid(sheet: dict, values) -> list[list]:
    """Computed values when the app sent them (same shape as rows), else the raw cells."""
    if isinstance(values, list) and values and all(isinstance(r, list) for r in values):
        return values
    return sheet["rows"]


def _pdf():
    from fpdf import FPDF

    from .analysis import _unicode_font

    pdf = FPDF()
    pdf.set_auto_page_break(True, 15)
    pdf.add_page()
    font = _unicode_font()
    if font:
        pdf.add_font("body", "", font)
        fam, enc = "body", (lambda s: s)
    else:
        fam, enc = "helvetica", (lambda s: s.encode("latin-1", "replace").decode("latin-1"))
    return pdf, fam, enc


def _pdf_table(pdf, fam, enc, header: list, rows: list[list]) -> None:
    n = max(1, len(header))
    w = (pdf.w - pdf.l_margin - pdf.r_margin) / n
    pdf.set_font(fam, size=10)
    for i, h in enumerate(header):
        pdf.cell(w, 8, enc(str(h))[:40], border=1, new_x="RIGHT", new_y="TOP")
    pdf.ln(8)
    pdf.set_font(fam, size=9)
    for r in rows:
        for i in range(n):
            v = r[i] if i < len(r) else ""
            pdf.cell(w, 7, enc(_fmt(v))[:40], border=1, new_x="RIGHT", new_y="TOP")
        pdf.ln(7)


def _fmt(v) -> str:
    if isinstance(v, float):
        return f"{v:,.2f}".rstrip("0").rstrip(".") if v != int(v) else f"{int(v):,}"
    return "" if v is None else str(v)


def _png_bytes(png: str) -> bytes:
    if not png:
        return b""
    m = re.match(r"^data:image/png;base64,(.+)$", png, re.S)
    if not m:
        raise HTTPException(400, "The chart picture must be a PNG.")
    data = base64.b64decode(m.group(1))
    if len(data) > 10 * 1024 * 1024:
        raise HTTPException(413, "That chart picture is too big.")
    return data


def _chart_png_server(c: dict) -> bytes:
    """A PNG drawn here (matplotlib) when the app didn't send one."""
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError as e:
        raise HTTPException(400, "Saving charts as pictures here needs the files add-on; use the PNG button in the app.") from e
    fig, ax = plt.subplots(figsize=(8, 4.5), dpi=150)
    if c["type"] == "pie" and c["series"]:
        ax.pie(c["series"][0]["values"], labels=c["labels"], autopct="%1.0f%%")
    else:
        n = max(1, len(c["series"]))
        for i, s in enumerate(c["series"]):
            xs = list(range(len(s["values"])))
            if c["type"] == "line":
                ax.plot(xs, s["values"], marker="o", label=s["name"])
            else:
                ax.bar([x + i * 0.8 / n for x in xs], s["values"], width=0.8 / n, label=s["name"])
        ax.set_xticks(range(len(c["labels"])), c["labels"])
        if n > 1:
            ax.legend()
    ax.set_title(c["title"])
    buf = io.BytesIO()
    fig.tight_layout()
    fig.savefig(buf, format="png")
    plt.close(fig)
    return buf.getvalue()


def export(kind: str, fmt: str, content, title: str = "", values=None, png: str = "") -> tuple[bytes, str, str]:
    """Returns (bytes, filename, mime type)."""
    if kind not in KINDS:
        raise HTTPException(400, "Unknown kind of live document.")
    fmt = fmt.lower().lstrip(".")
    if fmt not in FORMATS[kind]:
        raise HTTPException(400, f"A {kind} can be saved as {', '.join(FORMATS[kind])}.")
    try:
        if kind == "doc":
            data, title = _doc(fmt, str(content or ""), title)
        elif kind == "sheet":
            s = _sheet(content)
            title = title or s["title"]
            data = _sheet_export(fmt, s, _grid(s, values), title)
        else:
            c = _chart(content)
            title = title or c["title"]
            data = _chart_export(fmt, c, _png_bytes(png), title)
    except ImportError as e:
        raise HTTPException(400, f"Saving as .{fmt} needs the files add-on (Files page → Install).") from e
    except (ValueError, TypeError) as e:
        raise HTTPException(400, f"That {kind} could not be read: {e}") from e
    return data, f"{safe_name(title)}.{fmt}", MIME[fmt]


def _doc(fmt: str, md: str, title: str) -> tuple[bytes, str]:
    from .analysis import _blocks, _strip_md

    if not title:
        m = re.search(r"^#\s+(.+)$", md, re.M)
        title = m.group(1).strip() if m else "Document"
    if fmt == "md":
        return md.encode(), title
    if fmt == "docx":
        import docx

        d = docx.Document()
        for kind, level, text in _blocks(md):
            if kind == "h":
                d.add_heading(_strip_md(text), level - 1 if level > 1 or text != title else 0)
            elif kind == "li":
                d.add_paragraph(_strip_md(text), style="List Bullet")
            else:
                d.add_paragraph(_strip_md(text))
        buf = io.BytesIO()
        d.save(buf)
        return buf.getvalue(), title
    pdf, fam, enc = _pdf()
    for kind, level, text in _blocks(md):
        pdf.set_font(fam, size={1: 18, 2: 14, 3: 12}.get(level, 11))
        pdf.multi_cell(0, 6 if kind != "h" else 9, enc(("•  " if kind == "li" else "") + _strip_md(text)))
        pdf.ln(1)
    return bytes(pdf.output()), title


def _sheet_export(fmt: str, s: dict, grid: list[list], title: str) -> bytes:
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(s["columns"])
        for r in grid:
            w.writerow([_fmt(v) if isinstance(v, float) else v for v in r])
        return ("\ufeff" + buf.getvalue()).encode()  # BOM: Excel opens UTF-8 CSV correctly
    if fmt == "pdf":
        pdf, fam, enc = _pdf()
        pdf.set_font(fam, size=16)
        pdf.multi_cell(0, 9, enc(title))
        pdf.ln(2)
        _pdf_table(pdf, fam, enc, s["columns"], grid)
        return bytes(pdf.output())
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = safe_name(title)[:31]
    ws.append(s["columns"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="FFF3D6")
    for r in s["rows"]:
        ws.append(list(r))  # "=..." strings stay live Excel formulas
    for i, col in enumerate(s["columns"], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = min(40, max(10, len(col) + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _chart_export(fmt: str, c: dict, png: bytes, title: str) -> bytes:
    if fmt == "png":
        return png or _chart_png_server(c)
    header = ["", *[s["name"] for s in c["series"]]]
    rows = [[lab, *[(s["values"][i] if i < len(s["values"]) else 0) for s in c["series"]]] for i, lab in enumerate(c["labels"])]
    if fmt == "csv":
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(header)
        w.writerows(rows)
        return ("\ufeff" + buf.getvalue()).encode()
    if fmt == "pdf":
        pdf, fam, enc = _pdf()
        pdf.set_font(fam, size=16)
        pdf.multi_cell(0, 9, enc(title))
        img = png or _chart_png_server(c)
        pdf.image(io.BytesIO(img), w=pdf.w - pdf.l_margin - pdf.r_margin)
        pdf.ln(4)
        _pdf_table(pdf, fam, enc, header, rows)
        return bytes(pdf.output())
    import openpyxl
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = safe_name(title)[:31]
    ws.append(header)
    for r in rows:
        ws.append(r)
    ch = {"bar": BarChart, "line": LineChart, "pie": PieChart}[c["type"]]()
    ch.title = title
    n = len(rows)
    data = Reference(ws, min_col=2, max_col=1 + (1 if c["type"] == "pie" else len(c["series"])), min_row=1, max_row=n + 1)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n + 1))
    ws.add_chart(ch, f"{openpyxl.utils.get_column_letter(len(header) + 2)}2")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


REVISE_PROMPT = ("You edit one live {kind} for the user. Apply their instruction and return ONLY the full new {kind} {shape}. "
                 "Keep everything they didn't ask to change. No commentary.")
SHAPES = {"doc": "as Markdown", "sheet": 'as JSON {"title", "columns", "rows"} (formulas like "=SUM(B2:B5)" allowed; row 1 is the header)',
          "chart": 'as JSON {"title", "type": "bar"|"line"|"pie", "labels", "series": [{"name", "values"}]}'}


async def revise(kind: str, content, instruction: str, model: str | None = None) -> dict:
    """Ask the assistant to change a live block. Returns {"content"} in the same shape as was sent."""
    from . import llm

    if kind not in KINDS:
        raise HTTPException(400, "Unknown kind of live document.")
    if not (instruction or "").strip():
        raise HTTPException(400, "Say what to change.")
    current = content if isinstance(content, str) else json.dumps(content, ensure_ascii=False)
    msgs = [{"role": "system", "content": REVISE_PROMPT.format(kind=kind, shape=SHAPES[kind])},
            {"role": "user", "content": f"Current {kind}:\n{current[:60000]}\n\nInstruction: {instruction.strip()[:2000]}"}]
    try:
        text = await llm.complete(model, msgs, {"channel": "live_docs"})
    except HTTPException:
        raise
    except Exception as e:  # noqa: BLE001
        raise HTTPException(502, f"The model couldn't make that change: {str(e)[:200]}") from e
    text = re.sub(r"^\s*```[\w-]*\s*\n|\n```\s*$", "", (text or "").strip())
    if kind == "doc":
        if not text.strip():
            raise HTTPException(502, "The model sent back an empty document. Try again.")
        return {"content": text}
    d = llm.parse_json(text)
    if not isinstance(d, dict):
        raise HTTPException(502, "The model's answer wasn't a valid table or chart. Try saying it another way.")
    return {"content": _sheet(d) if kind == "sheet" else _chart(d)}
