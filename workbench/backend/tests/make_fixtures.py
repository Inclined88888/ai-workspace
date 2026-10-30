"""Builds one small file of every kind the Files page reads, so tests need no binary fixtures in git."""

import io
import json
import struct
import wave
import zipfile
from email.message import EmailMessage
from pathlib import Path


def build(d: Path) -> dict[str, Path]:
    d.mkdir(parents=True, exist_ok=True)
    out = {}

    def w(name, data):
        p = d / name
        p.write_bytes(data if isinstance(data, bytes) else data.encode())
        out[name] = p
        return p

    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Item", "Category", "Cost", "Date"])
    for i, (it, cat, c) in enumerate([("Timber", "Materials", 1200), ("Labour", "Labour", 3400), ("Paint", "Materials", 300),
                                     ("Scaffold", "Hire", 850), ("Labour", "Labour", 3400)], start=2):
        ws.append([it, cat, c, f"2026-10-{i:02d}"])
    ws["C8"] = "=SUM(C2:C6)"
    ws["B8"] = "Total"
    wb.save(d / "budget.xlsx")
    out["budget.xlsx"] = d / "budget.xlsx"
    w("sales.csv", "region;month;units;revenue\n" + "".join(f"{r};2026-{m:02d};{u};{u * 12.5}\n" for r in ("North", "South") for m, u in ((1, 10), (2, 14), (3, 9))))

    import docx

    for name, pay, term in (("contract_v1.docx", "30 days", "12 months"), ("contract_v2.docx", "14 days", "24 months")):
        doc = docx.Document()
        doc.add_heading("Services Agreement", 1)
        doc.add_paragraph("1. The Supplier will provide maintenance services to the Client.")
        doc.add_paragraph(f"2. Payment. The Client must pay each invoice within {pay} of receipt.")
        doc.add_paragraph(f"3. Term. This agreement runs for {term} and renews automatically unless either party gives notice.")
        doc.add_paragraph("4. Liability. The Supplier's total liability is limited to the fees paid in the previous year.")
        doc.add_paragraph("5. Termination. Either party may terminate for material breach on 30 days' written notice.")
        doc.add_paragraph("6. Indemnity. The Client shall indemnify the Supplier against third party claims. Governing law: Queensland.")
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "Service", "Fee", "Maintenance", "$500" if pay == "30 days" else "$650"
        doc.save(d / name)
        out[name] = d / name

    from fpdf import FPDF

    pdf = FPDF()
    for n in (1, 2):
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.multi_cell(0, 8, f"Inspection report page {n}.\n\nThe slab on level {n} passed inspection on 3 November 2026.\n\nDefects: none recorded.")
    w("report.pdf", bytes(pdf.output()))

    from pptx import Presentation

    pr = Presentation()
    s = pr.slides.add_slide(pr.slide_layouts[1])
    s.shapes.title.text = "Quarterly update"
    s.placeholders[1].text = "Revenue up 12%\nTwo new clients"
    s.notes_slide.notes_text_frame.text = "Mention the hiring plan"
    pr.save(d / "update.pptx")
    out["update.pptx"] = d / "update.pptx"

    w("notes.md", "# Meeting notes\n\nDiscussed the roster.\n\n## Actions\n\n- Alex to order pipe\n")
    w("data.json", json.dumps({"project": "Tower B", "tasks": [{"id": i, "name": f"Task {i}", "done": i % 2 == 0} for i in range(5)]}))
    w("orders.xml", "<orders>" + "".join(f"<order id='{i}'><sku>A{i}</sku><qty>{i}</qty></order>" for i in range(4)) + "</orders>")
    w("analysis.ipynb", json.dumps({"cells": [{"cell_type": "markdown", "source": ["# Growth model"]},
                                              {"cell_type": "code", "source": ["import numpy as np\n", "x = np.arange(10)"], "outputs": []}],
                                    "metadata": {"kernelspec": {"name": "python3"}}, "nbformat": 4, "nbformat_minor": 5}))
    w("paper.tex", r"\documentclass{article}\begin{document}\section{Introduction}We solve $x^2-5x+6=0$." "\n"
                   r"\begin{equation}E=mc^2\end{equation}\section{Method}Text.\end{document}")
    w("script.py", "import os\n\ndef total(xs):\n    return sum(xs)\n\nclass Ledger:\n    pass\n")

    import pandas as pd

    pd.DataFrame({"temp": [21.5, 22.1, 19.8], "site": ["A", "B", "C"]}).to_parquet(d / "readings.parquet")
    out["readings.parquet"] = d / "readings.parquet"
    import h5py
    import numpy as np

    with h5py.File(d / "experiment.h5", "w") as h:
        h.create_dataset("run1/signal", data=np.linspace(0, 1, 100))
        h.attrs["instrument"] = "Spectrometer"
    out["experiment.h5"] = d / "experiment.h5"
    cards = [f"{k:<8}= {v:>20}" for k, v in (("SIMPLE", "T"), ("BITPIX", "16"), ("NAXIS", "2"), ("NAXIS1", "4"), ("NAXIS2", "4"), ("OBJECT", "'M31     '"))] + ["END"]
    w("galaxy.fits", "".join(c.ljust(80) for c in cards).ljust(2880).encode() + b"\0" * 2880)

    from PIL import Image

    img = Image.new("RGB", (64, 32), (240, 200, 80))
    img.save(d / "photo.png")
    out["photo.png"] = d / "photo.png"
    buf = io.BytesIO()
    with wave.open(buf, "wb") as wv:
        wv.setnchannels(1)
        wv.setsampwidth(2)
        wv.setframerate(8000)
        wv.writeframes(struct.pack("<h", 0) * 8000)
    w("voice.wav", buf.getvalue())

    import ezdxf

    dx = ezdxf.new()
    dx.layers.add("WALLS")
    msp = dx.modelspace()
    msp.add_line((0, 0), (10, 0), dxfattribs={"layer": "WALLS"})
    msp.add_text("KITCHEN", dxfattribs={"layer": "0"})
    dx.saveas(d / "plan.dxf")
    out["plan.dxf"] = d / "plan.dxf"

    m = EmailMessage()
    m["From"], m["To"], m["Subject"] = "site@builder.example", "me@example.com", "Variation 4 approved"
    m.set_content("Hi, variation 4 is approved. Costs attached.")
    m.add_attachment(b"item,cost\nExtra window,900\n", maintype="text", subtype="csv", filename="variation.csv")
    w("mail.eml", m.as_bytes())

    zb = io.BytesIO()
    with zipfile.ZipFile(zb, "w") as z:
        z.write(d / "notes.md", "docs/notes.md")
        z.write(d / "sales.csv", "docs/sales.csv")
        z.writestr("../evil.txt", "nope")
    w("bundle.zip", zb.getvalue())
    w("mystery.bin", b"\x00\x01binary-blob PART-NUMBER-XJ9 more\x00\xff" * 10)
    return out
